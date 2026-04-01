from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(r"D:\Acode\Android\complete\MedicalSystemDemo")
DOC_DIR = ROOT / "交付" / "doc"


def drawio_file(title, vertices, edges):
    def xml_text(value):
        base = escape(value)
        encoded = []
        for ch in base:
            if ord(ch) > 127:
                encoded.append(f"&#{ord(ch)};")
            else:
                encoded.append(ch)
        return "".join(encoded)

    cells = ['<mxCell id="0"/>', '<mxCell id="1" parent="0"/>']
    for v in vertices:
        cells.append(
            f'<mxCell id="{v["id"]}" value="{xml_text(v["text"])}" style="{v["style"]}" vertex="1" parent="1">'
            f'<mxGeometry x="{v["x"]}" y="{v["y"]}" width="{v["w"]}" height="{v["h"]}" as="geometry"/>'
            f"</mxCell>"
        )
    for e in edges:
        cells.append(
            f'<mxCell id="{e["id"]}" value="{xml_text(e.get("text", ""))}" style="{e["style"]}" edge="1" parent="1" source="{e["source"]}" target="{e["target"]}">'
            f'<mxGeometry relative="1" as="geometry"/>'
            f"</mxCell>"
        )
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<mxfile host="app.diagrams.net" modified="2026-04-01T00:00:00.000Z" agent="Codex" version="24.7.17">\n'
        f'  <diagram id="diagram-1" name="{xml_text(title)}">\n'
        '    <mxGraphModel dx="1400" dy="900" grid="1" gridSize="10" guides="1" tooltips="1" connect="1" arrows="1" fold="1" page="1" pageScale="1" pageWidth="1654" pageHeight="1169" math="0" shadow="0">\n'
        "      <root>\n"
        "        "
        + "\n        ".join(cells)
        + "\n"
        "      </root>\n"
        "    </mxGraphModel>\n"
        "  </diagram>\n"
        "</mxfile>\n"
    )


BOX = "rounded=1;whiteSpace=wrap;html=1;fillColor=#dae8fc;strokeColor=#6c8ebf;fontSize=14;"
PROCESS = "rounded=1;whiteSpace=wrap;html=1;fillColor=#d5e8d4;strokeColor=#82b366;fontSize=14;"
DECISION = "rhombus;whiteSpace=wrap;html=1;fillColor=#ffe6cc;strokeColor=#d79b00;fontSize=14;"
DATA = "shape=parallelogram;perimeter=parallelogramPerimeter;whiteSpace=wrap;html=1;fillColor=#f8cecc;strokeColor=#b85450;fontSize=14;"
EDGE = "endArrow=block;html=1;rounded=1;strokeWidth=1.5;"


def write_drawio_files():
    arch_vertices = [
        dict(id="v1", text="控制台宿主程序\nMedicalSystemDemo_app.cpp", x=60, y=220, w=220, h=80, style=BOX),
        dict(id="v2", text="EDL 桥接层\nMedicalSystemDemo.edl\nMedicalSystemDemo_u/t", x=340, y=220, w=220, h=80, style=BOX),
        dict(id="v3", text="Enclave 业务层\nMedicalSystemDemo.cpp", x=620, y=220, w=220, h=80, style=PROCESS),
        dict(id="v4", text="内存态系统状态\nMsdSystemState", x=900, y=120, w=220, h=80, style=PROCESS),
        dict(id="v5", text="权限与会话控制\ncurrent_user / has_role", x=900, y=220, w=220, h=80, style=PROCESS),
        dict(id="v6", text="患者档案与病历数组\nusers / patients / records", x=900, y=320, w=220, h=80, style=PROCESS),
        dict(id="v7", text="OCall 外部服务\n打印日志 / 文件读写 / 获取时间", x=620, y=420, w=220, h=80, style=BOX),
        dict(id="v8", text="密封状态文件\nsealed_system_state.bin", x=900, y=420, w=220, h=80, style=DATA),
    ]
    arch_edges = [
        dict(id="e1", source="v1", target="v2", style=EDGE, text="ECall"),
        dict(id="e2", source="v2", target="v3", style=EDGE, text="受信接口"),
        dict(id="e3", source="v3", target="v4", style=EDGE, text="维护"),
        dict(id="e4", source="v3", target="v5", style=EDGE, text="执行"),
        dict(id="e5", source="v3", target="v6", style=EDGE, text="读写"),
        dict(id="e6", source="v3", target="v7", style=EDGE, text="OCall"),
        dict(id="e7", source="v7", target="v8", style=EDGE, text="保存/读取"),
    ]

    login_vertices = [
        dict(id="v1", text="开始", x=80, y=40, w=120, h=50, style=PROCESS),
        dict(id="v2", text="宿主程序采集用户名和密码", x=60, y=130, w=220, h=60, style=BOX),
        dict(id="v3", text="调用 ecall_login_user", x=70, y=230, w=200, h=60, style=PROCESS),
        dict(id="v4", text="系统已初始化？", x=90, y=330, w=160, h=80, style=DECISION),
        dict(id="v5", text="存在活动会话？", x=320, y=330, w=160, h=80, style=DECISION),
        dict(id="v6", text="查找用户\nfind_user_index", x=550, y=330, w=180, h=60, style=PROCESS),
        dict(id="v7", text="密码匹配？", x=570, y=430, w=160, h=80, style=DECISION),
        dict(id="v8", text="写入 active_session_index\n返回角色", x=550, y=550, w=200, h=70, style=PROCESS),
        dict(id="v9", text="返回错误码\n未初始化/重复登录/认证失败", x=300, y=550, w=200, h=80, style=DATA),
        dict(id="v10", text="宿主程序切换到对应角色菜单", x=540, y=670, w=220, h=60, style=BOX),
    ]
    login_edges = [
        dict(id="e1", source="v1", target="v2", style=EDGE),
        dict(id="e2", source="v2", target="v3", style=EDGE),
        dict(id="e3", source="v3", target="v4", style=EDGE),
        dict(id="e4", source="v4", target="v9", style=EDGE, text="否"),
        dict(id="e5", source="v4", target="v5", style=EDGE, text="是"),
        dict(id="e6", source="v5", target="v9", style=EDGE, text="是"),
        dict(id="e7", source="v5", target="v6", style=EDGE, text="否"),
        dict(id="e8", source="v6", target="v7", style=EDGE),
        dict(id="e9", source="v7", target="v9", style=EDGE, text="否"),
        dict(id="e10", source="v7", target="v8", style=EDGE, text="是"),
        dict(id="e11", source="v8", target="v10", style=EDGE),
    ]

    patient_vertices = [
        dict(id="v1", text="开始", x=80, y=40, w=120, h=50, style=PROCESS),
        dict(id="v2", text="管理员/患者发起档案新增或修改", x=50, y=120, w=240, h=60, style=BOX),
        dict(id="v3", text="ecall_upsert_patient_profile", x=70, y=220, w=200, h=60, style=PROCESS),
        dict(id="v4", text="输入是否合法？", x=90, y=320, w=160, h=80, style=DECISION),
        dict(id="v5", text="对应患者账号是否存在？", x=320, y=320, w=180, h=80, style=DECISION),
        dict(id="v6", text="是否有编辑权限？\n管理员或本人", x=560, y=320, w=180, h=80, style=DECISION),
        dict(id="v7", text="查找档案槽位\nfind_patient_index", x=800, y=320, w=180, h=60, style=PROCESS),
        dict(id="v8", text="不存在则分配新槽位\nallocate_patient_slot", x=790, y=430, w=200, h=70, style=PROCESS),
        dict(id="v9", text="写入 full_name / id_card / phone / address", x=780, y=540, w=220, h=70, style=PROCESS),
        dict(id="v10", text="宿主侧 flush_state\n触发密封保存", x=790, y=660, w=200, h=60, style=DATA),
        dict(id="v11", text="返回错误码", x=360, y=540, w=160, h=70, style=DATA),
    ]
    patient_edges = [
        dict(id="e1", source="v1", target="v2", style=EDGE),
        dict(id="e2", source="v2", target="v3", style=EDGE),
        dict(id="e3", source="v3", target="v4", style=EDGE),
        dict(id="e4", source="v4", target="v11", style=EDGE, text="否"),
        dict(id="e5", source="v4", target="v5", style=EDGE, text="是"),
        dict(id="e6", source="v5", target="v11", style=EDGE, text="否"),
        dict(id="e7", source="v5", target="v6", style=EDGE, text="是"),
        dict(id="e8", source="v6", target="v11", style=EDGE, text="否"),
        dict(id="e9", source="v6", target="v7", style=EDGE, text="是"),
        dict(id="e10", source="v7", target="v8", style=EDGE),
        dict(id="e11", source="v8", target="v9", style=EDGE),
        dict(id="e12", source="v9", target="v10", style=EDGE),
    ]

    record_vertices = [
        dict(id="v1", text="开始", x=80, y=40, w=120, h=50, style=PROCESS),
        dict(id="v2", text="医生发起新增/修改/删除病历", x=50, y=120, w=240, h=60, style=BOX),
        dict(id="v3", text="ecall_create_record\n或 ecall_update_record\n或 ecall_delete_record", x=40, y=220, w=260, h=90, style=PROCESS),
        dict(id="v4", text="当前用户是否医生？", x=90, y=350, w=160, h=80, style=DECISION),
        dict(id="v5", text="患者档案/病历是否存在？", x=320, y=350, w=180, h=80, style=DECISION),
        dict(id="v6", text="若为修改或删除\n是否为病历创建医生？", x=560, y=350, w=200, h=80, style=DECISION),
        dict(id="v7", text="新增：分配 record_id\n写入 doctor_username 和 created_at", x=820, y=320, w=230, h=80, style=PROCESS),
        dict(id="v8", text="修改：覆盖 diagnosis / prescription / note", x=820, y=430, w=230, h=70, style=PROCESS),
        dict(id="v9", text="删除：memset 清空记录槽位", x=820, y=530, w=230, h=60, style=PROCESS),
        dict(id="v10", text="宿主侧 flush_state\n密封并写入 sealed_system_state.bin", x=830, y=640, w=220, h=70, style=DATA),
        dict(id="v11", text="返回错误码", x=360, y=560, w=160, h=70, style=DATA),
    ]
    record_edges = [
        dict(id="e1", source="v1", target="v2", style=EDGE),
        dict(id="e2", source="v2", target="v3", style=EDGE),
        dict(id="e3", source="v3", target="v4", style=EDGE),
        dict(id="e4", source="v4", target="v11", style=EDGE, text="否"),
        dict(id="e5", source="v4", target="v5", style=EDGE, text="是"),
        dict(id="e6", source="v5", target="v11", style=EDGE, text="否"),
        dict(id="e7", source="v5", target="v6", style=EDGE, text="是"),
        dict(id="e8", source="v6", target="v11", style=EDGE, text="否"),
        dict(id="e9", source="v6", target="v7", style=EDGE, text="新增"),
        dict(id="e10", source="v6", target="v8", style=EDGE, text="修改"),
        dict(id="e11", source="v6", target="v9", style=EDGE, text="删除"),
        dict(id="e12", source="v7", target="v10", style=EDGE),
        dict(id="e13", source="v8", target="v10", style=EDGE),
        dict(id="e14", source="v9", target="v10", style=EDGE),
    ]

    files = {
        "系统总体架构.drawio": drawio_file("系统总体架构", arch_vertices, arch_edges),
        "登录与权限控制流程.drawio": drawio_file("登录与权限控制流程", login_vertices, login_edges),
        "患者档案CRUD流程.drawio": drawio_file("患者档案CRUD流程", patient_vertices, patient_edges),
        "病历CRUD与密封存储流程.drawio": drawio_file("病历CRUD与密封存储流程", record_vertices, record_edges),
    }

    for name, content in files.items():
        (DOC_DIR / name).write_text(content, encoding="utf-8")
    return list(files.keys())


def write_markdown(drawio_names):
    lines = [
        "# 项目细节补充说明",
        "",
        "## 1. 问题回应",
        "老师指出的问题是“目前整体偏文字描述，缺少架构、接口、数据结构、实现流程和代码支撑”。本说明专门补足这些细节，用于论文正文、项目报告和答辩说明。",
        "",
        "## 2. 系统总体架构",
        "系统采用“宿主程序 + EDL 桥接 + Enclave 业务逻辑 + 外部密封文件”四层结构。",
        "",
        "1. 宿主程序层：`MedicalSystemDemo_app.cpp`，负责控制台交互、菜单分发、ECall 发起、OCall 实现、文件路径解析。",
        "2. 接口桥接层：`MedicalSystemDemo.edl` 定义 Enclave 可调用接口和 OCall 接口，再由 `sgx_edger8r` 生成 `MedicalSystemDemo_u.*`、`MedicalSystemDemo_t.*`。",
        "3. Enclave 业务层：`MedicalSystemDemo.cpp`，负责权限校验、用户/档案/病历管理、系统状态密封与解封。",
        "4. 外部存储层：宿主程序通过 `ocall_save_blob` 和 `ocall_load_blob` 将密封后的二进制数据写入 `sealed_system_state.bin`。",
        "",
        "配套 draw.io 文件：`系统总体架构.drawio`。",
        "",
        "## 3. 模块划分与职责",
        "",
        "| 模块 | 主要文件 | 主要职责 |",
        "|---|---|---|",
        "| 宿主层 | `MedicalSystemDemo_app.cpp` | 用户交互、菜单控制、ECall 调用、OCall 文件读写、日志输出 |",
        "| 接口层 | `MedicalSystemDemo.edl` | 定义可信接口和非可信接口边界 |",
        "| 可信业务层 | `MedicalSystemDemo.cpp` | 用户认证、权限控制、患者档案 CRUD、病历 CRUD、状态持久化 |",
        "| 共享数据定义 | `shared_types.h` | 枚举、结构体、容量上限、错误码定义 |",
        "",
        "## 4. 核心数据结构设计",
        "",
        "本系统没有直接引入数据库，而是采用 Enclave 内部固定容量顺序表。原因是毕业设计第一阶段的重点在于“可信计算 + 密封存储 + 业务原型”，因此用静态数组结构更容易说明可信边界和访问控制逻辑。",
        "",
        "| 结构体 | 关键字段 | 作用 |",
        "|---|---|---|",
        "| `MsdUserAccount` | `username`、`password`、`role`、`active` | 存储管理员/医生/患者账号信息 |",
        "| `MsdPatientProfile` | `patient_username`、`full_name`、`id_card`、`phone`、`address`、`active` | 存储患者基本档案 |",
        "| `MsdMedicalRecord` | `record_id`、`patient_username`、`doctor_username`、`diagnosis`、`prescription`、`note`、`created_at`、`active` | 存储病历信息 |",
        "| `MsdSystemState` | `initialized`、`next_record_id`、`active_session_index`、`users[]`、`patients[]`、`records[]` | 整个系统在 Enclave 内部的运行状态 |",
        "",
        "- 最多 32 个用户：`MSD_MAX_USERS = 32`",
        "- 最多 32 个患者档案：`MSD_MAX_PATIENTS = 32`",
        "- 最多 128 条病历：`MSD_MAX_RECORDS = 128`",
        "",
        "```cpp",
        "typedef struct MsdSystemState {",
        "    uint32_t magic;",
        "    uint32_t version;",
        "    int initialized;",
        "    int next_record_id;",
        "    int active_session_index;",
        "    MsdUserAccount users[MSD_MAX_USERS];",
        "    MsdPatientProfile patients[MSD_MAX_PATIENTS];",
        "    MsdMedicalRecord records[MSD_MAX_RECORDS];",
        "} MsdSystemState;",
        "```",
        "",
        "## 5. 接口设计",
        "",
        "系统已定义系统管理、用户管理、患者档案管理和病历管理四类 ECall 接口，以及日志、文件和时间三类 OCall 接口。",
        "",
        "## 6. 权限控制实现",
        "",
        "系统采用“单活动会话 + 角色校验 + 对象级权限判断”三层控制。",
        "",
        "- `current_user()` 根据 `active_session_index` 返回当前登录用户。",
        "- `has_role()` 检查是否为管理员、医生或患者。",
        "- `can_view_patient()` 与 `can_edit_patient()` 进一步判断患者对象级权限。",
        "- 系统只允许一个活动会话，重复登录会返回 `MSD_ERR_SESSION_ACTIVE`。",
        "",
        "配套 draw.io 文件：`登录与权限控制流程.drawio`。",
        "",
        "## 7. 患者档案 CRUD 实现",
        "",
        "新增/修改通过 `ecall_upsert_patient_profile` 实现，读取通过 `ecall_get_patient_profile` 与 `ecall_list_patients` 实现，删除通过 `ecall_delete_patient_profile` 实现。",
        "",
        "关键限制是：只有管理员可以删除患者档案，而且有病历关联时不能删除。",
        "",
        "```cpp",
        "if (patient_has_records(patient_username)) {",
        "    write_result(result, MSD_ERR_PROFILE_HAS_RECORDS);",
        "    return;",
        "}",
        "memset(&g_state.patients[patient_index], 0, sizeof(g_state.patients[patient_index]));",
        "```",
        "",
        "配套 draw.io 文件：`患者档案CRUD流程.drawio`。",
        "",
        "## 8. 病历 CRUD 实现",
        "",
        "新增病历由 `ecall_create_record` 实现，内部自动分配 `record_id`、记录创建医生并写入时间戳；修改和删除分别由 `ecall_update_record`、`ecall_delete_record` 实现。",
        "",
        "```cpp",
        "record->record_id = g_state.next_record_id++;",
        "copy_text(record->doctor_username, sizeof(record->doctor_username), user->username);",
        "ocall_get_time(record->created_at, sizeof(record->created_at));",
        "```",
        "",
        "病历修改和删除都要求“当前医生必须是病历创建者”。",
        "",
        "配套 draw.io 文件：`病历CRUD与密封存储流程.drawio`。",
        "",
        "## 9. 密封存储与恢复实现",
        "",
        "系统通过 `sgx_seal_data` 和 `sgx_unseal_data` 对 `MsdSystemState` 整体密封和恢复，宿主程序通过 `ocall_save_blob`、`ocall_load_blob` 与外部文件交互。",
        "",
        "## 10. 可直接用于论文的结论",
        "",
        "本项目已经形成了明确的可信边界、明确的数据结构、明确的接口集合和明确的 CRUD 实现路径，不再只是概念描述，而是一个可运行、可说明、可扩展的可信医疗原型系统。",
        "",
        "## 11. 配套 draw.io 文件",
        "",
    ]
    for name in drawio_names:
        lines.append(f"- `{name}`")
    (DOC_DIR / "项目细节补充说明.md").write_text("\n".join(lines), encoding="utf-8")


def set_cn_font(run, name, size=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def build_doc(drawio_names):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(11)

    for sname in ["Heading 1", "Heading 2", "Heading 3"]:
        s = doc.styles[sname]
        s.font.name = "黑体"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r = p.add_run("项目细节补充说明")
    set_cn_font(r, "黑体", 16, True)

    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r = p.add_run("用于回应“缺少架构、接口、数据结构、实现细节和流程图支撑”的问题")
    set_cn_font(r, "宋体", 11)

    def para(text, code=False):
        p = doc.add_paragraph(style="Normal")
        r = p.add_run(text)
        if code:
            set_cn_font(r, "Consolas", 10.5)
        else:
            set_cn_font(r, "宋体", 11)

    def code_block(text):
        for line in text.strip("\n").splitlines():
            para(line, code=True)

    doc.add_heading("1. 问题回应", level=1)
    para("老师指出的问题是“目前整体偏文字描述，缺少架构、接口、数据结构、实现流程和代码支撑”。本说明专门补足这些细节，用于论文正文、项目报告和答辩说明。")

    doc.add_heading("2. 系统总体架构", level=1)
    para("系统采用“宿主程序 + EDL 桥接 + Enclave 业务逻辑 + 外部密封文件”四层结构。")
    for t in [
        "宿主程序层：MedicalSystemDemo_app.cpp，负责控制台交互、菜单分发、ECall 发起、OCall 实现、文件路径解析。",
        "接口桥接层：MedicalSystemDemo.edl 定义 Enclave 可调用接口和 OCall 接口，再由 sgx_edger8r 生成 MedicalSystemDemo_u.*、MedicalSystemDemo_t.*。",
        "Enclave 业务层：MedicalSystemDemo.cpp，负责权限校验、用户/档案/病历管理、系统状态密封与解封。",
        "外部存储层：宿主程序通过 ocall_save_blob 和 ocall_load_blob 将密封后的二进制数据写入 sealed_system_state.bin。",
    ]:
        para("• " + t)
    para("配套 draw.io 文件：系统总体架构.drawio。")

    doc.add_heading("3. 模块划分与职责", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["模块", "主要文件", "主要职责"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in [
        ("宿主层", "MedicalSystemDemo_app.cpp", "用户交互、菜单控制、ECall 调用、OCall 文件读写、日志输出"),
        ("接口层", "MedicalSystemDemo.edl", "定义可信接口和非可信接口边界"),
        ("可信业务层", "MedicalSystemDemo.cpp", "用户认证、权限控制、患者档案 CRUD、病历 CRUD、状态持久化"),
        ("共享数据定义", "shared_types.h", "枚举、结构体、容量上限、错误码定义"),
    ]:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v

    doc.add_heading("4. 核心数据结构设计", level=1)
    para("本系统没有直接引入数据库，而是采用 Enclave 内部固定容量顺序表。原因是毕业设计第一阶段的重点在于“可信计算 + 密封存储 + 业务原型”，因此用静态数组结构更容易说明可信边界和访问控制逻辑。")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["结构体", "关键字段", "作用"]):
        table.rows[0].cells[i].text = h
    for row in [
        ("MsdUserAccount", "username、password、role、active", "存储管理员/医生/患者账号信息"),
        ("MsdPatientProfile", "patient_username、full_name、id_card、phone、address、active", "存储患者基本档案"),
        ("MsdMedicalRecord", "record_id、patient_username、doctor_username、diagnosis、prescription、note、created_at、active", "存储病历信息"),
        ("MsdSystemState", "initialized、next_record_id、active_session_index、users[]、patients[]、records[]", "整个系统在 Enclave 内部的运行状态"),
    ]:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    for t in [
        "当前容量限制：最多 32 个用户、32 个患者档案、128 条病历。",
        "查找通过顺序扫描实现，逻辑简单，便于论文说明。",
        "删除操作通过 memset 清空槽位，并利用 active 作为逻辑删除标记。",
        "next_record_id 保证病历编号单调递增。",
    ]:
        para("• " + t)
    code_block(
        """
typedef struct MsdSystemState {
    uint32_t magic;
    uint32_t version;
    int initialized;
    int next_record_id;
    int active_session_index;
    MsdUserAccount users[MSD_MAX_USERS];
    MsdPatientProfile patients[MSD_MAX_PATIENTS];
    MsdMedicalRecord records[MSD_MAX_RECORDS];
} MsdSystemState;
"""
    )

    doc.add_heading("5. 接口设计", level=1)
    para("系统已定义系统管理、用户管理、患者档案管理和病历管理四类 ECall 接口，以及日志、文件和时间三类 OCall 接口。")

    doc.add_heading("6. 权限控制实现", level=1)
    for t in [
        "系统采用“单活动会话 + 角色校验 + 对象级权限判断”三层控制。",
        "current_user() 根据 active_session_index 返回当前登录用户。",
        "has_role() 检查是否为管理员、医生或患者。",
        "can_view_patient() 与 can_edit_patient() 进一步判断患者对象级权限。",
        "管理员可以注册账号、查看全部用户、查看和维护所有患者档案。",
        "医生可以查看患者列表、查看患者档案、创建/修改/删除自己创建的病历。",
        "患者只能查看和修改自己的档案，只能查看自己的病历。",
        "系统只允许一个活动会话，重复登录会返回 MSD_ERR_SESSION_ACTIVE。",
    ]:
        para("• " + t)
    para("配套 draw.io 文件：登录与权限控制流程.drawio。")

    doc.add_heading("7. 患者档案 CRUD 实现", level=1)
    for t in [
        "新增/修改通过 ecall_upsert_patient_profile 实现，读取通过 ecall_get_patient_profile 与 ecall_list_patients 实现，删除通过 ecall_delete_patient_profile 实现。",
        "新增或修改前会校验：系统是否初始化、输入是否合法、患者账号是否存在、当前用户是否有编辑权限。",
        "删除前会校验：当前用户是否为管理员、患者是否仍然关联病历。",
    ]:
        para("• " + t)
    code_block(
        """
if (patient_has_records(patient_username)) {
    write_result(result, MSD_ERR_PROFILE_HAS_RECORDS);
    return;
}
memset(&g_state.patients[patient_index], 0, sizeof(g_state.patients[patient_index]));
"""
    )
    para("配套 draw.io 文件：患者档案CRUD流程.drawio。")

    doc.add_heading("8. 病历 CRUD 实现", level=1)
    for t in [
        "新增病历由 ecall_create_record 实现，内部自动分配 record_id、记录创建医生并写入时间戳。",
        "修改和删除分别由 ecall_update_record、ecall_delete_record 实现。",
        "病历修改和删除都要求“当前医生必须是病历创建者”。",
    ]:
        para("• " + t)
    code_block(
        """
record->record_id = g_state.next_record_id++;
copy_text(record->doctor_username, sizeof(record->doctor_username), user->username);
ocall_get_time(record->created_at, sizeof(record->created_at));
"""
    )
    para("配套 draw.io 文件：病历CRUD与密封存储流程.drawio。")

    doc.add_heading("9. 密封存储与恢复实现", level=1)
    for t in [
        "系统通过 sgx_seal_data 和 sgx_unseal_data 对 MsdSystemState 整体密封和恢复，宿主程序通过 ocall_save_blob、ocall_load_blob 与外部文件交互。",
        "保存时会先复制状态快照，并将 active_session_index 重置为 -1，避免把登录态直接持久化。",
        "恢复时会校验 magic 和 version，防止状态文件损坏或版本不匹配。",
    ]:
        para("• " + t)
    for line in [
        "[App] 已从密态状态文件读取: D:\\...\\sealed_system_state.bin",
        "[App] 密态状态文件已保存到: D:\\...\\sealed_system_state.bin",
        "[Enclave-Log] [Enclave] 系统状态已密封并请求写入文件: sealed_system_state.bin",
    ]:
        para(line, code=True)

    doc.add_heading("10. 宿主程序菜单如何驱动业务", level=1)
    for t in [
        "guest_menu：初始化系统、加载系统状态、登录。",
        "admin_menu：注册账号、维护患者档案、查看列表、删除档案、保存状态。",
        "doctor_menu：查看患者、查看病历、新增/修改/删除病历、保存状态。",
        "patient_menu：查看/修改本人档案、查看本人病历、保存状态。",
        "main() 根据 g_current_role 切换菜单，因此宿主层只负责流程编排，核心权限判断全部在 Enclave 内实现。",
    ]:
        para("• " + t)

    doc.add_heading("11. 当前设计的工程意义与局限", level=1)
    for t in [
        "意义：把用户认证、患者档案和病历信息都放在 Enclave 内部状态中管理，通过 SGX sealing 实现密态持久化，通过 EDL 明确划分可信与非可信边界。",
        "局限：当前仍是固定容量数组原型，不是完整数据库；没有 Remote Attestation、审计日志落盘、医生授权范围模型；界面仍为控制台原型。",
    ]:
        para("• " + t)

    doc.add_heading("12. 可直接用于论文的结论", level=1)
    for t in [
        "本项目已经形成了明确的可信边界、明确的数据结构、明确的接口集合和明确的 CRUD 实现路径，不再只是概念描述，而是一个可运行、可说明、可扩展的可信医疗原型系统。",
        "系统内部通过 MsdSystemState 统一管理用户、患者档案和病历数据，采用固定容量顺序表作为基础数据结构，并通过 find_*、allocate_*、append_format 等函数实现查找、分配和结果输出。",
        "在业务实现上，患者档案和病历均已实现增删查改，其中新增/修改操作在完成后会触发系统状态密封保存；删除操作则附带权限和一致性约束，例如“有病历关联的患者档案不可删除”“医生只能修改或删除自己创建的病历”。",
    ]:
        para(t)

    para("附：本说明配套 draw.io XML 文件如下，可直接在 draw.io 中选择“文件 -> 导入”打开。")
    for name in drawio_names:
        para("• " + name, code=True)

    tmp = ROOT / "Project_Detail_Supplement.docx"
    doc.save(str(tmp))
    final = DOC_DIR / "项目细节补充说明.docx"
    if final.exists():
        final.unlink()
    tmp.replace(final)


def main():
    drawio_names = write_drawio_files()
    write_markdown(drawio_names)
    build_doc(drawio_names)


if __name__ == "__main__":
    main()
